from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas import PriceInput
from app.services.price_import import inspect_prices, preview_prices, normalize_header

FIXTURES = Path(__file__).parent / "fixtures" / "price_import"


def make_book(headers, rows, title=None):
    book = Workbook()
    sheet = book.active
    if title:
        sheet.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers))
        )
        sheet.cell(1, 1, title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    "fixture, expected_header, expected_name",
    [
        ("canonical_first_sheet.xlsx", 0, "Modul fotovoltaik 550 Wp"),
        ("albanian_merged_title.xlsx", 2, "Çelik për armim"),
        ("english_reordered.xlsx", 0, "Furnizim dhe montim konstruksioni FV"),
        ("legacy_albanian.xls", 2, "Çelik për armim"),
    ],
)
def test_real_workbook_layouts(fixture, expected_header, expected_name):
    data = (FIXTURES / fixture).read_bytes()
    inspected = inspect_prices(data, fixture)
    assert inspected["sheets"][0]["header_row"] == expected_header
    result = preview_prices(data, fixture)
    assert result["error_count"] == 0
    assert result["valid_count"] > 0
    assert result["rows"][0]["values"]["name"] == expected_name
    assert (
        result["sheet"] == inspected["default_sheet"] == inspected["sheets"][0]["name"]
    )


def test_first_worksheet_used_even_when_another_is_active():
    data = (FIXTURES / "canonical_first_sheet.xlsx").read_bytes()
    info = inspect_prices(data, "file.xlsx")
    assert info["default_sheet"] == "Çmimet"
    assert preview_prices(data, "file.xlsx")["rows"][0]["values"]["code"] == "PV-01"
    other = preview_prices(data, "file.xlsx", sheet="Other active sheet")
    assert other["sheet"] == "Other active sheet"
    with pytest.raises(ValueError, match="Fleta"):
        preview_prices(data, "file.xlsx", sheet="nonexistent")


@pytest.mark.parametrize(
    "code,name,description,unit,price",
    [
        ("CODE", "NAME", "DESCRIPTION", "UNIT", "PRICE"),
        (" KODI ", " EMËRTIMI ", "PËRSHKRIM", " NJËSIA ", " ÇMIMI_NJËSI "),
        ("kod", "artikulli", "pershkrim", "njesi", "cmimi"),
        ("nr", "zëri", "description", "njësi", "unit_price"),
        ("no", "pershkrimi", "description", "unit", "price"),
        ("code", "përshkrimi", "description", "njesia", "çmimi"),
    ],
)
def test_user_header_aliases_case_order_and_whitespace(
    code, name, description, unit, price
):
    data = make_book(
        [price, unit, name, code, description],
        [["12,50", "m³", "Emërtim", "abc-01", "Përshkrimi teknik"]],
    )
    result = preview_prices(data, "catalog.xlsx")
    assert result["error_count"] == 0
    values = result["rows"][0]["values"]
    assert values["code"] == "ABC-01"
    assert values["name"] == "Emërtim"
    assert values["description"] == "Përshkrimi teknik"
    assert values["price"] == "12.50"


def test_bom_unicode_spaces_accents_and_underscores():
    assert normalize_header("\ufeff  ÇMIMI_\u00a0NJËSI\u2009 ") == "cmimi njesi"
    data = make_book(
        ["\ufeffnAmE\u00a0", "\u2009UNIT\u202f", " UNIT_Price "],
        [["Çelik", "kg", "1\u202f234,56"]],
    )
    result = preview_prices(data, "x.xlsx", currency="EUR")
    values = result["rows"][0]["values"]
    assert values["description"] == "Çelik"
    assert values["price"] == "1234.56"
    assert values["currency"] == "EUR"
    assert values["code"] == "IMPORT-000002"
    assert result["rows"][0]["warnings"]


def test_export_canonical_optional_fields_are_retained():
    original = PriceInput(
        code="A",
        name="Beton",
        description="Beton C25/30",
        unit="m³",
        price="125.123456",
        currency="EUR",
        synonyms="betoni",
        technical_specs="C25/30",
        category="Beton",
        subcategory="Furnizim",
        price_kind="purchase",
        supply_included=True,
        installation_included=False,
        supplier="Furnitori",
        price_date="2026-01-15",
        valid_until="2026-12-31",
        vat_included=True,
        vat_rate="20.0000",
        material_cost="100.001",
        labor_cost="2",
        machinery_cost="3",
        transport_cost="4",
        other_cost="5",
        notes="Shënim",
        active=False,
        reason="Ofertë",
    )
    fields = list(PriceInput.model_fields)
    data = make_book(
        [*fields, "Ignored Vendor Field"],
        [[*[original.model_dump()[field] for field in fields], "ignored"]],
    )
    result = preview_prices(data, "prices.xlsx")
    assert result["error_count"] == 0
    values = PriceInput.model_validate(result["rows"][0]["values"])
    assert values == original
    assert "Ignored Vendor Field" not in result["mapping"]


def test_description_only_fallback_and_numeric_code():
    data = make_book(["Nr", "Description", "Unit", "Price"], [[7, "Kabllo", "m", 0]])
    result = preview_prices(data, "catalog.xlsx")
    row = result["rows"][0]
    assert row["status"] == "valid"
    assert row["values"]["code"] == "7"
    assert row["values"]["name"] == row["values"]["description"] == "Kabllo"
    assert Decimal(row["values"]["price"]) == Decimal(
        "0"
    )  # Explicit zero is permitted; missing is not.


def test_ambiguous_aliases_require_manual_mapping():
    data = (FIXTURES / "ambiguous_prices.xlsx").read_bytes()
    sheet = inspect_prices(data, "prices.xlsx")["sheets"][0]
    assert sheet["ambiguous_headers"] == {"price": [2, 3]}
    assert "price" not in sheet["suggested_mapping"]
    result = preview_prices(data, "prices.xlsx")
    assert result["valid_count"] == 0
    assert result["error_count"] == 1
    assert any(error["field"] == "price" for error in result["rows"][0]["errors"])
    manual = preview_prices(
        data, "prices.xlsx", mapping={"name": 0, "unit": 1, "price": 2}, header_row=0
    )
    assert manual["error_count"] == 0
    assert manual["rows"][0]["values"]["price"] == "1.20"


def test_manual_arbitrary_columns_and_generated_code_source_row():
    data = make_book(
        ["Supplier article", "Cost 2026", "UM", "Unused"],
        [["Panel FV", "25,99", "copë", "extra"]],
        title="Katalogu i furnitorit",
    )
    result = preview_prices(
        data,
        "custom.xlsx",
        mapping={"description": 0, "price": 1, "unit": 2},
        header_row=1,
        code_prefix="FV26",
    )
    assert result["valid_count"] == 1
    row = result["rows"][0]
    assert row["source_row"] == 3
    assert row["values"]["code"] == "FV26-000003"


def test_unknown_headers_and_explicit_empty_mapping_report_errors_without_throwing():
    data = make_book(["X", "Y", "Z"], [["Beton", "m³", 100]])
    for mapping in (None, {}):
        result = preview_prices(data, "unknown.xlsx", mapping=mapping)
        assert result["valid_count"] == 0
        assert result["error_count"] == 1
        assert {"name", "unit", "price"} <= {
            error["field"] for error in result["rows"][0]["errors"]
        }


@pytest.mark.parametrize(
    "price",
    [
        None,
        "",
        "NaN",
        "Infinity",
        "-1",
        True,
        "0.1234567",
        "12.1234560",
        "1 23 4",
        "12,34,56",
        "1,23.56",
        "€ 12,00",
    ],
)
def test_invalid_prices_are_row_specific_without_raw_value_leaks(price):
    data = make_book(["Name", "Unit", "Price"], [["Beton", "m³", price]])
    result = preview_prices(data, "invalid.xlsx")
    assert result["valid_count"] == 0
    assert result["error_count"] == 1
    row = result["rows"][0]
    assert row["source_row"] == 2
    assert row["values"] == {}
    assert any(error["field"] == "price" for error in row["errors"])
    assert all(set(error) == {"field", "message"} for error in row["errors"])


@pytest.mark.parametrize(
    "text,expected",
    [
        ("1.234,56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("1 234,56", "1234.56"),
        ("1\u00a0234,56", "1234.56"),
        ("12,50", "12.50"),
        ("0.123456", "0.123456"),
        ("1,234,567", "1234567"),
        ("1.234.567", "1234567"),
    ],
)
def test_decimal_locales(text, expected):
    result = preview_prices(
        make_book(["Name", "Unit", "Price"], [["A", "kg", text]]), "x.xlsx"
    )
    assert result["error_count"] == 0
    assert Decimal(result["rows"][0]["values"]["price"]) == Decimal(expected)


def test_skipped_rows_and_missing_units():
    data = make_book(
        ["Code", "Name", "Unit", "Price"],
        [
            [None, None, None, None],
            [None, "Kapitulli I", None, None],
            [None, "Nëntotal", None, 25],
            ["X", "Pa njësi", None, "20"],
            ["Y", "Pa çmim", "m", None],
            ["Z", "Artikull", "m", "25"],
        ],
    )
    result = preview_prices(data, "rows.xlsx")
    assert (result["valid_count"], result["error_count"], result["skipped_count"]) == (
        1,
        2,
        3,
    )
    assert result["rows"][-1]["source_row"] == 7
    assert result["rows"][-1]["values"]["code"] == "Z"


def test_duplicate_codes_case_normalized():
    data = make_book(
        ["Code", "Name", "Unit", "Price"],
        [["Aa", "One", "m", 1], ["aa", "Two", "m", 2]],
    )
    result = preview_prices(data, "duplicate.xlsx")
    assert result["valid_count"] == 1
    assert result["error_count"] == 1
    assert result["rows"][1]["errors"][0]["field"] == "code"


def test_dates_booleans_and_optional_values():
    data = make_book(
        [
            "name",
            "unit",
            "price",
            "price_date",
            "valid_until",
            "vat_included",
            "active",
        ],
        [["A", "kg", "20", datetime(2026, 1, 15), date(2027, 1, 15), "Po", "Jo"]],
    )
    result = preview_prices(data, "dates.xlsx")
    values = result["rows"][0]["values"]
    assert result["error_count"] == 0
    assert values["price_date"] == "2026-01-15"
    assert values["valid_until"] == "2027-01-15"
    assert values["vat_included"] is True
    assert values["active"] is False


def test_detect_headers_after_merged_title_in_first_fifty_rows():
    book = Workbook()
    sheet = book.active
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Titulli i katalogut"
    sheet.append([])
    sheet.cell(49, 1, "Emërtimi")
    sheet.cell(49, 2, "Njësia")
    sheet.cell(49, 3, "Çmimi")
    sheet.append(["Artikulli", "m", "12,00"])
    output = BytesIO()
    book.save(output)
    info = inspect_prices(output.getvalue(), "header.xlsx")["sheets"][0]
    assert info["header_row"] == 48
    result = preview_prices(output.getvalue(), "header.xlsx")
    assert result["valid_count"] == 1
    assert result["rows"][0]["source_row"] == 50


def test_formatted_blank_tail_does_not_use_import_data_row_limit():
    book = Workbook()
    sheet = book.active
    sheet.append(["Name", "Unit", "Price"])
    sheet.append(["A", "kg", "10"])
    sheet.cell(20000, 3).font = Font(bold=True)
    output = BytesIO()
    book.save(output)
    result = preview_prices(output.getvalue(), "formatted.xlsx")
    assert result["valid_count"] == 1
    assert result["error_count"] == 0
    assert result["skipped_count"] == 19998


@pytest.mark.parametrize(
    "mapping",
    [
        {"name": 0, "unit": 1, "price": True},
        {"name": 0, "unit": 1, "price": 99},
        {"name": 0, "unit": 1, "price": 1},
        {"owner_id": 0},
    ],
)
def test_invalid_mappings_rejected(mapping):
    with pytest.raises(ValueError):
        preview_prices(
            make_book(["Name", "Unit", "Price"], [["A", "kg", 1]]),
            "x.xlsx",
            mapping=mapping,
        )


def test_invalid_filename_file_types_and_security_bounds(monkeypatch):
    data = (FIXTURES / "canonical_first_sheet.xlsx").read_bytes()
    for name, contents in [
        ("file.csv", data),
        ("file.xls", data),
        ("file.xlsx", b"not excel"),
        ("file.xls", b"bad OLE"),
        ("file.xlsx", b""),
    ]:
        with pytest.raises(ValueError):
            inspect_prices(contents, name)
    legacy = (FIXTURES / "legacy_albanian.xls").read_bytes()
    import app.services.price_import as service

    monkeypatch.setattr(service, "MAX_COLUMNS", 2)
    with pytest.raises(ValueError, match="kufirin"):
        inspect_prices(legacy, "file.xls")
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("../escaped", "no")
        archive.writestr("xl/workbook.xml", "no")
    with pytest.raises(ValueError):
        inspect_prices(output.getvalue(), "file.xlsx")


def test_formula_price_is_not_executed_or_assigned():
    data = make_book(["Name", "Unit", "Price"], [["A", "kg", "=1+2"]])
    result = preview_prices(data, "formula.xlsx")
    assert result["valid_count"] == 0
    assert result["error_count"] == 1
    assert result["rows"][0]["values"] == {}


@pytest.mark.parametrize(
    "value, review",
    [
        ("1,234", True),
        ("1.234", True),
        ("0.123", False),
        (1.234, False),
        ("1234.567", False),
    ],
)
def test_ambiguous_textual_separators_require_explicit_preview_review(value, review):
    data = make_book(["Name", "Unit", "Price"], [["A", "kg", value]])
    result = preview_prices(data, "ambiguous.xlsx")
    assert result["valid_count"] == 1
    row = result["rows"][0]
    assert row["numeric_review_required"] is review
    assert any("mijëshe" in message for message in row["warnings"]) is review


def test_ambiguous_optional_cost_also_requires_numeric_review():
    data = make_book(
        ["Name", "Unit", "Price", "material_cost"], [["A", "kg", "10", "1.234"]]
    )
    row = preview_prices(data, "cost.xlsx")["rows"][0]
    assert row["status"] == "valid"
    assert row["numeric_review_required"] is True
    assert any(message.startswith("material_cost:") for message in row["warnings"])


def test_lazy_xml_parse_errors_become_controlled_workbook_errors():
    # Corrupt data beyond the initial XML chunk: openpyxl loads dimensions first,
    # and encounters this error only while consuming the read-only worksheet.
    data = make_book(
        ["Name", "Unit", "Price"],
        [[f"Artikull {index}", "kg", 10] for index in range(1500)],
    )
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"</sheetData>", b"</brokenSheetData>")
            target.writestr(member.filename, content)
    with pytest.raises(ValueError, match="dëmtuar|pasigurt"):
        preview_prices(output.getvalue(), "corrupt.xlsx")


def test_nonfinite_legacy_number_preview_is_json_serializable(monkeypatch):
    import json
    import app.services.price_import as service

    sheet = service._Sheet(
        "Sheet",
        2,
        3,
        lambda: iter([["Name", "Unit", "Price"], ["A", "kg", float("nan")]]),
    )
    monkeypatch.setattr(
        service,
        "_open_catalog",
        lambda data, filename: service._Book([sheet], lambda: None),
    )
    info = inspect_prices(b"dummy", "x.xls")
    json.dumps(info, allow_nan=False)
    assert info["sheets"][0]["preview"][1][2] == "nan"
    assert preview_prices(b"dummy", "x.xls")["error_count"] == 1


@pytest.mark.parametrize("header", ["Description", "Emërtimi"])
def test_long_single_source_description_retained_with_short_display_name(header):
    description = (
        "Furnizim dhe montim i konstruksionit metalik për panele fotovoltaike. " * 12
    ).strip()
    data = make_book([header, "Unit", "Price"], [[description, "kg", "12,50"]])
    result = preview_prices(data, "long-description.xlsx")
    assert result["valid_count"] == 1
    row = result["rows"][0]
    assert row["values"]["description"] == description
    assert row["values"]["name"] == description[:300].rstrip()
    assert any("300" in message and "plotë" in message for message in row["warnings"])


def test_explicit_separate_oversized_name_is_not_silently_truncated():
    data = make_book(
        ["Name", "Description", "Unit", "Price"],
        [["N" * 301, "Përshkrim i veçantë", "kg", "12,50"]],
    )
    result = preview_prices(data, "explicit-name.xlsx")
    assert result["valid_count"] == 0
    assert result["error_count"] == 1
    assert any(error["field"] == "name" for error in result["rows"][0]["errors"])
