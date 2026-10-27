"""Regression tests for the interactive price importer, using actual workbooks."""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from openpyxl import Workbook, load_workbook
from app.db import get_db
from app.main import app
from app.models import PriceImportPreview


def workbook(headers, rows, title=False, active_second=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Preventivi ndërtimit"
    if title:
        sheet.merge_cells("A1:F1")
        sheet["A1"] = "PREVENTIV – Punime ndërtimore"
        sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if active_second:
        other = book.create_sheet("Shënime")
        other.append(["Ky nuk është preventivi"])
        book.active = 1
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def preview(client, data, filename="preventiv.xlsx", **options):
    return client.post(
        "/api/prices/import/preview", files={"file": (filename, data)}, data=options
    )


def test_albanian_title_rows_first_sheet_and_atomic_confirmation(client, user):
    data = workbook(
        [" NJËSIA ", " ÇMIMI_NJËSI ", " EMËRTIMI ", " NR "],
        [
            ["m³", "1.234,56", "Beton C25/30", "B-01"],
            [None, None, None, None],
            [None, None, "TOTALI", None],
        ],
        title=True,
        active_second=True,
    )
    response = preview(client, data, currency="EUR")
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["sheet"] == "Preventivi ndërtimit"
    assert result["header_row"] == 2
    assert result["valid_count"] == 1
    assert result["can_commit"] is True
    assert result["rows"][0]["values"]["description"] == "Beton C25/30"
    assert result["rows"][0]["values"]["price"] == "1234.56"
    assert result["rows"][0]["values"]["currency"] == "EUR"
    assert client.get("/api/prices").json() == []
    assert client.get("/api/prices/history").json() == []
    token = result["preview_id"]
    assert (
        client.post(
            "/api/prices/import/commit", json={"preview_id": token, "confirm": False}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/prices/import/commit", json={"preview_id": token, "confirm": True}
        ).json()["imported"]
        == 1
    )
    saved = client.get("/api/prices").json()
    assert saved[0]["code"] == "B-01"
    assert saved[0]["price"] == "1234.560000"
    assert (
        client.post(
            "/api/prices/import/commit", json={"preview_id": token, "confirm": True}
        ).status_code
        == 404
    )


def test_invalid_prices_row_specific_and_nothing_saved(client, user):
    data = workbook(
        ["name", "unit", "price"],
        [
            ["Beton", "m³", "10"],
            ["Zhavorr", "m³", "bad"],
            ["Punë", "orë", -1],
            ["Çelik", "kg", "NaN"],
            ["Rërë", "m³", None],
            ["PV", "copë", "0.0000001"],
        ],
    )
    result = preview(client, data).json()
    assert result["valid_count"] == 1
    assert result["error_count"] == 5
    invalid = [r for r in result["rows"] if r["status"] == "invalid"]
    assert [r["source_row"] for r in invalid] == [3, 4, 5, 6, 7]
    assert all(any(e["field"] == "price" for e in r["errors"]) for r in invalid)
    assert result["can_commit"] is False
    assert (
        client.post(
            "/api/prices/import/commit",
            json={"preview_id": result["preview_id"], "confirm": True},
        ).status_code
        == 422
    )
    assert client.get("/api/prices").json() == []
    assert client.get("/api/prices/history").json() == []


def test_unknown_columns_manual_mapping_and_missing_code_description(client, user):
    data = workbook(["A", "B", "C"], [["10,50", "copë", "Modul fotovoltaik"]])
    initial = preview(client, data)
    assert initial.status_code == 200, initial.text
    assert initial.json()["can_commit"] is False
    remapped = preview(
        client,
        data,
        mapping='{"name":2,"price":0,"unit":1}',
        header_row="0",
        code_prefix="FV",
        currency="EUR",
    )
    assert remapped.status_code == 200, remapped.text
    result = remapped.json()
    assert result["can_commit"]
    assert result["rows"][0]["values"]["code"].startswith("FV-")
    assert result["rows"][0]["values"]["name"] == "Modul fotovoltaik"
    assert result["rows"][0]["values"]["description"] == "Modul fotovoltaik"
    assert result["rows"][0]["warnings"]
    assert (
        client.post(
            "/api/prices/import/commit",
            json={"preview_id": result["preview_id"], "confirm": True},
        ).status_code
        == 200
    )


def test_duplicate_code_preview_and_stale_commit_are_safe(client, price):
    data = workbook(["code", "name", "unit", "price"], [["C25", "Beton", "m³", "10"]])
    duplicate = preview(client, data).json()
    assert not duplicate["can_commit"]
    assert duplicate["rows"][0]["errors"][0]["field"] == "code"
    data = workbook(["code", "name", "unit", "price"], [["NEW", "Beton", "m³", "10"]])
    initial = preview(client, data).json()
    assert initial["can_commit"]
    client.post(
        "/api/prices",
        json={
            "code": "NEW",
            "name": "Other",
            "description": "Other",
            "unit": "m³",
            "price": "20",
        },
    )
    assert (
        client.post(
            "/api/prices/import/commit",
            json={"preview_id": initial["preview_id"], "confirm": True},
        ).status_code
        == 422
    )
    assert len(client.get("/api/prices").json()) == 2


def test_preview_ownership_expiry_and_validation(client, user):
    data = workbook(["name", "unit", "price"], [["A", "m", "1"]])
    result = preview(client, data).json()
    token = result["preview_id"]
    with next(app.dependency_overrides[get_db]()) as db:
        entry = db.get(PriceImportPreview, token)
        entry.expires_at = datetime.now(timezone.utc) - timedelta(minutes=1)
        db.commit()
    assert (
        client.post(
            "/api/prices/import/commit", json={"preview_id": token, "confirm": True}
        ).status_code
        == 410
    )
    token = preview(client, data).json()["preview_id"]
    client.post("/api/auth/logout")
    assert preview(client, data).status_code == 401
    client.post(
        "/api/auth/register",
        json={"email": "outsider@example.com", "password": "long-password-123"},
    )
    assert (
        client.post(
            "/api/prices/import/commit", json={"preview_id": token, "confirm": True}
        ).status_code
        == 404
    )
    assert preview(client, data, mapping="[]").status_code == 422
    assert preview(client, data, filename="file.csv").status_code == 422
    assert preview(client, b"not excel").status_code == 422
    assert preview(client, data, header_row="-1").status_code == 422


def test_export_format_preserved_and_english_headers_case_order(client, user):
    data = workbook(
        ["  PRICE ", " NAME ", " UNIT ", " CODE "], [["0", "Test", "kg", "test"]]
    )
    result = preview(client, data).json()
    assert result["valid_count"] == 1
    assert (
        client.post(
            "/api/prices/import/commit",
            json={"preview_id": result["preview_id"], "confirm": True},
        ).status_code
        == 200
    )
    exported = client.get("/api/prices/export")
    book = load_workbook(BytesIO(exported.content))
    assert [c.value for c in book.active[1]][:5] == [
        "code",
        "name",
        "description",
        "synonyms",
        "technical_specs",
    ]
    assert book.active["A2"].value == "TEST"


def test_ambiguous_text_requires_interactive_review_not_silent_legacy_save(
    client, user
):
    data = workbook(["code", "name", "unit", "price"], [["A", "Beton", "m³", "1,234"]])
    direct = client.post("/api/prices/import", files={"file": ("boq.xlsx", data)})
    assert direct.status_code == 422
    assert "paraprake" in direct.json()["detail"]
    assert client.get("/api/prices").json() == []
    proposed = preview(client, data).json()
    assert proposed["rows"][0]["numeric_review_required"]
    assert proposed["rows"][0]["warnings"]
    assert proposed["rows"][0]["values"]["price"] == "1.234"
    assert (
        client.post(
            "/api/prices/import/commit",
            json={"preview_id": proposed["preview_id"], "confirm": True},
        ).status_code
        == 200
    )
