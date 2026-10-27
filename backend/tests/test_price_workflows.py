"""API integration coverage for atomic price import and resource snapshots."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook


def price_body(code, **changes):
    return {
        "code": code,
        "name": code,
        "description": f"Përshkrimi {code}",
        "unit": "kg",
        "price": "10",
        "currency": "ALL",
        **changes,
    }


def workbook_bytes(headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def upload(client, data, filename="prices.xlsx"):
    return client.post(
        "/api/prices/import",
        files={
            "file": (
                filename,
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def create(client, code, **changes):
    response = client.post("/api/prices", json=price_body(code, **changes))
    assert response.status_code == 201, response.text
    return response.json()


def test_prices_export_and_import_round_trip(client, user):
    original = create(
        client,
        "PV-01",
        description="Modul fotovoltaik 550 Wp",
        price="123.456789",
        currency="EUR",
        supply_included=True,
        synonyms="FV, panel",
        technical_specs="550 Wp",
        vat_included=True,
        vat_rate="20",
    )
    response = client.get("/api/prices/export")
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith('filename="prices.xlsx"')
    book = load_workbook(BytesIO(response.content), data_only=False)
    assert book.active["A2"].value == "PV-01"
    assert "price" in [cell.value for cell in book.active[1]]
    # Import as another user: shared file contents never carry database ownership.
    client.post("/api/auth/logout")
    client.post(
        "/api/auth/register",
        json={"email": "importer@example.com", "password": "password-for-importer"},
    )
    imported = upload(client, response.content)
    assert imported.status_code == 200, imported.text
    assert imported.json() == {"imported": 1}
    rows = client.get("/api/prices").json()
    assert len(rows) == 1
    assert rows[0]["owner_id"] != original["owner_id"]
    for key in (
        "code",
        "description",
        "price",
        "currency",
        "supply_included",
        "vat_included",
        "synonyms",
        "technical_specs",
    ):
        assert rows[0][key] == original[key]


def test_price_export_prevents_formula_injection(client, user):
    create(client, "X", description='=HYPERLINK("https://evil.example")')
    response = client.get("/api/prices/export")
    book = load_workbook(BytesIO(response.content), data_only=False)
    column = [cell.value for cell in book.active[1]].index("description") + 1
    cell = book.active.cell(2, column)
    assert cell.data_type == "s"
    assert cell.value.startswith("'=HYPERLINK")


def test_import_validation_is_atomic(client, user):
    data = workbook_bytes(
        ["code", "name", "description", "unit", "price"],
        [
            ["A", "Valid", "Valid row", "kg", "10"],
            ["B", "Invalid", "Invalid price", "kg", "not-a-number"],
        ],
    )
    response = upload(client, data)
    assert response.status_code == 422
    assert client.get("/api/prices").json() == []
    assert client.get("/api/prices/history").json() == []


def test_duplicate_import_rolls_back_earlier_rows(client, user):
    create(client, "EXISTING")
    data = workbook_bytes(
        ["code", "name", "description", "unit", "price"],
        [
            ["NEW", "New", "Should roll back", "kg", "10"],
            ["EXISTING", "Duplicate", "Already exists", "kg", "20"],
        ],
    )
    response = upload(client, data)
    assert response.status_code == 409
    assert [row["code"] for row in client.get("/api/prices").json()] == ["EXISTING"]
    assert len(client.get("/api/prices/history").json()) == 1


def test_reject_duplicate_import_headers_and_legacy_excel(client, user):
    data = workbook_bytes(
        ["code", "name", "description", "unit", "price", "price"],
        [["A", "A", "A", "kg", "1", "2"]],
    )
    assert upload(client, data).status_code == 422
    assert upload(client, data, filename="prices.xls").status_code == 422
    assert upload(client, b"garbage").status_code == 422
    assert client.get("/api/prices").json() == []


def test_composite_net_cost_waste_and_version_snapshots(client, user):
    labor = create(
        client, "LABOR", unit="orë", price="120", vat_included=True, vat_rate="20"
    )
    steel = create(client, "STEEL", unit="kg", price="30")
    work = create(client, "WORK", unit="m²", price="999")
    response = client.put(
        f"/api/prices/{work['id']}/resources",
        json={
            "reason": "Analiza e parë",
            "resources": [
                {"item_id": labor["id"], "coefficient": "2", "waste_percent": "0"},
                {"item_id": steel["id"], "coefficient": "3", "waste_percent": "10"},
            ],
        },
    )
    assert response.status_code == 200, response.text
    version = response.json()
    assert Decimal(version["direct_cost"]) == Decimal("299")  # 2*100 net + 3*30*1.10
    assert version["number"] == 1
    assert version["resources"][0]["unit"] == "orë"
    assert version["resources"][0]["original_price"] == "120.000000"
    assert version["resources"][0]["vat_included"] is True
    current = next(
        row for row in client.get("/api/prices").json() if row["id"] == work["id"]
    )
    assert Decimal(current["price"]) == Decimal("299")
    assert current["vat_included"] is False
    assert current["price_kind"] == "direct"
    # Source updates do not retroactively alter an analysis or its parent price.
    result = client.put(
        f"/api/prices/{steel['id']}", json=price_body("STEEL", price="60")
    )
    assert result.status_code == 200
    versions = client.get(f"/api/prices/{work['id']}/resources").json()
    assert Decimal(versions[0]["direct_cost"]) == Decimal("299")
    response = client.put(
        f"/api/prices/{work['id']}/resources",
        json={
            "reason": "Koeficient i rishikuar",
            "resources": [{"item_id": steel["id"], "coefficient": "2"}],
        },
    )
    assert response.status_code == 200
    assert response.json()["number"] == 2
    assert Decimal(response.json()["direct_cost"]) == Decimal("120")
    history = client.get("/api/prices/history", params={"item_id": work["id"]}).json()
    assert Decimal(history[0]["old_price"]) == Decimal("299")
    assert Decimal(history[0]["new_price"]) == Decimal("120")


def test_composite_currency_requires_explicit_exchange_rate(client, user):
    resource = create(client, "EUR", currency="EUR", price="2")
    work = create(client, "ALL", currency="ALL")
    body = {
        "reason": "Burime",
        "resources": [{"item_id": resource["id"], "coefficient": "3"}],
    }
    response = client.put(f"/api/prices/{work['id']}/resources", json=body)
    assert response.status_code == 422
    assert "EUR:ALL" in response.json()["detail"]
    assert client.get(f"/api/prices/{work['id']}/resources").json() == []
    response = client.put("/api/settings", json={"exchange_rates": {"EUR:ALL": "100"}})
    assert response.status_code == 200, response.text
    response = client.put(f"/api/prices/{work['id']}/resources", json=body)
    assert response.status_code == 200, response.text
    assert Decimal(response.json()["direct_cost"]) == Decimal("600")
    assert response.json()["resources"][0]["original_currency"] == "EUR"
    assert response.json()["resources"][0]["converted_currency"] == "ALL"


def test_composite_invalid_empty_self_archived_and_overflow(client, user):
    resource = create(client, "RESOURCE", price="1000000000000")
    work = create(client, "WORK")
    endpoint = f"/api/prices/{work['id']}/resources"
    assert (
        client.put(endpoint, json={"reason": "Empty", "resources": []}).status_code
        == 422
    )
    assert (
        client.put(
            endpoint,
            json={
                "reason": "Self",
                "resources": [{"item_id": work["id"], "coefficient": "1"}],
            },
        ).status_code
        == 422
    )
    assert (
        client.put(
            endpoint,
            json={
                "reason": "Overflow",
                "resources": [{"item_id": resource["id"], "coefficient": "1000000000"}],
            },
        ).status_code
        == 422
    )
    assert client.delete(f"/api/prices/{resource['id']}").status_code == 204
    assert (
        client.put(
            endpoint,
            json={
                "reason": "Archived",
                "resources": [{"item_id": resource["id"], "coefficient": "1"}],
            },
        ).status_code
        == 422
    )
    assert client.get(endpoint).json() == []


def test_composite_authorization_for_parent_and_resource(client, user):
    private_resource = create(client, "PRIVATE")
    private_work = create(client, "PRIVATE-WORK")
    client.post("/api/auth/logout")
    assert (
        client.post(
            "/api/auth/register",
            json={
                "email": "other-resource@example.com",
                "password": "another-secure-password",
            },
        ).status_code
        == 201
    )
    own_work = create(client, "OWN-WORK")
    body = {
        "reason": "Analysis",
        "resources": [{"item_id": private_resource["id"], "coefficient": "1"}],
    }
    assert (
        client.put(f"/api/prices/{own_work['id']}/resources", json=body).status_code
        == 404
    )
    assert (
        client.put(f"/api/prices/{private_work['id']}/resources", json=body).status_code
        == 404
    )
    assert client.get(f"/api/prices/{private_work['id']}/resources").status_code == 404
    assert (
        client.get(
            "/api/prices/history", params={"item_id": private_resource["id"]}
        ).json()
        == []
    )
