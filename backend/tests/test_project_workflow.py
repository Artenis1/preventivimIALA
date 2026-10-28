from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook, load_workbook


def upload_boq(
    client,
    project,
    description="Furnizim beton C25/30",
    unit="m³",
    code="C25",
    quantity="2",
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Preventivi"
    ws.append(["Nr", "Kodi", "Përshkrimi", "Njësia", "Sasia", "Çmimi", "Vlera"])
    ws.append(["1.01", code, description, unit, quantity, 999, 1998])
    ws.append(["", "", "TOTAL", "", "", "", 1998])
    stream = BytesIO()
    wb.save(stream)
    response = client.post(
        f"/api/projects/{project['id']}/imports",
        files={"file": ("boq.xlsx", stream.getvalue())},
    )
    assert response.status_code == 200, response.text
    info = response.json()
    sheet = info["sheets"][0]
    response = client.post(
        f"/api/projects/{project['id']}/imports/{info['upload_id']}/commit",
        json={
            "sheet": sheet["name"],
            "header_row": sheet["header_row"],
            "mapping": sheet["suggested_mapping"],
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1
    return response.json(), info


def create_project(client):
    response = client.post(
        "/api/projects",
        json={
            "name": "Parku FV",
            "currency": "EUR",
            "overhead_percent": "10",
            "profit_percent": "5",
            "vat_rate": "20",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_full_workflow_and_snapshots(client, price):
    project = create_project(client)
    pid = project["id"]
    imported, info = upload_boq(client, project)
    row = imported["rows"][0]
    assert row["unit_price"] is None
    assert imported["totals"]["incomplete"]
    result = client.post(f"/api/projects/{pid}/match")
    assert result.status_code == 200, result.text
    match = result.json()["rows"][0]["match_result"]
    assert match["method"] == "code"
    preview = client.get(f"/api/projects/{pid}/approve-preview").json()
    assert preview["rows"] == []
    assert result.json()["rows"][0]["decision"] == "approved"
    assert result.json()["totals"]["grand_total"] == "276.00"
    assert (
        client.post(f"/api/projects/{pid}/finalize", json={"confirm": True}).status_code
        == 200
    )
    body = {
        k: price[k]
        for k in ("code", "name", "description", "unit", "price", "currency")
    }
    body["price"] = "200"
    assert client.put(f"/api/prices/{price['id']}", json=body).status_code == 200
    frozen = client.get(f"/api/projects/{pid}").json()
    assert Decimal(frozen["rows"][0]["unit_price"]) == Decimal("100")
    assert frozen["versions"][0]["snapshot"]["totals"]["grand_total"] == "276.00"
    assert (
        client.patch(
            f"/api/projects/{pid}/rows/{row['id']}",
            json={
                "unit_price": "300",
                "decision": "approved",
                "confirm_override": True,
            },
        ).status_code
        == 409
    )
    for format in ("standard", "original"):
        export = client.get(f"/api/projects/{pid}/export?format={format}")
        assert export.status_code == 200, export.text
        wb = load_workbook(BytesIO(export.content))
        assert len(wb.worksheets) >= 2
    client.post(f"/api/projects/{pid}/reopen", json={"confirm": True})
    differences = client.post(
        f"/api/projects/{pid}/refresh-prices", json={"confirm": False}
    ).json()
    assert differences["applied"] is False
    assert Decimal(differences["differences"][0]["new_price"]) == 200
    assert (
        Decimal(client.get(f"/api/projects/{pid}").json()["rows"][0]["unit_price"])
        == 100
    )
    assert (
        client.post(
            f"/api/projects/{pid}/refresh-prices", json={"confirm": True}
        ).status_code
        == 200
    )
    assert (
        client.post(f"/api/projects/{pid}/finalize", json={"confirm": True}).status_code
        == 200
    )
    compare = client.get(f"/api/projects/{pid}/compare?left=1&right=2").json()
    assert Decimal(compare["grand_total_difference"]) == 276
    assert client.get("/api/audit").json()


def test_missing_currency_and_incompatible_units(client, price):
    project = create_project(client)
    imported, _ = upload_boq(client, project, unit="copë")
    row = imported["rows"][0]
    pid = project["id"]
    match = client.post(f"/api/projects/{pid}/match").json()["rows"][0]["match_result"]
    assert match["classification"] != "high"
    assert match["suggested_unit_price"] is None
    assert (
        client.patch(
            f"/api/projects/{pid}/rows/{row['id']}",
            json={"item_id": price["id"], "decision": "approved", "confirm_override": True},
        ).status_code
        == 200
    )
    assert (
        client.post(f"/api/projects/{pid}/finalize", json={"confirm": True}).status_code
        == 409
    )


def test_decision_preserved_and_revised(client, price):
    project = create_project(client)
    pid = project["id"]
    imported, _ = upload_boq(client, project)
    row = imported["rows"][0]
    url = f"/api/projects/{pid}/rows/{row['id']}"
    assert (
        client.patch(
            url,
            json={"item_id": price["id"], "decision": "approved", "unit_price": "125"},
        ).status_code
        == 200
    )
    client.post(f"/api/projects/{pid}/match")
    assert (
        Decimal(client.get(f"/api/projects/{pid}").json()["rows"][0]["unit_price"])
        == 125
    )
    assert client.patch(url, json={"decision": "rejected"}).status_code == 409
    assert (
        client.patch(
            url, json={"decision": "rejected", "confirm_override": True}
        ).status_code
        == 200
    )
    client.post(f"/api/projects/{pid}/match")
    assert (
        client.get(f"/api/projects/{pid}").json()["rows"][0]["decision"] == "rejected"
    )
    assert (
        client.patch(
            url,
            json={"decision": "approved", "unit_price": "77", "confirm_override": True},
        ).status_code
        == 200
    )


def test_project_isolation_upload_validation_and_duplicate_commit(client, price):
    project = create_project(client)
    pid = project["id"]
    imported, info = upload_boq(client, project)
    assert (
        client.post(
            f"/api/projects/{pid}/imports/{info['upload_id']}/commit",
            json={
                "sheet": "Preventivi",
                "mapping": {"description": 2, "unit": 3, "quantity": 4},
            },
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/api/projects/{pid}/imports",
            files={"file": ("../../bad.xlsx", b"not a workbook")},
        ).status_code
        == 422
    )
    assert (
        client.post(
            f"/api/projects/{pid}/imports", files={"file": ("bad.xls", b"bad")}
        ).status_code
        == 422
    )
    client.post("/api/auth/logout")
    client.post(
        "/api/auth/register",
        json={"email": "private@example.com", "password": "safe-password-1234"},
    )
    assert client.get(f"/api/projects/{pid}").status_code == 404
    assert client.post(f"/api/projects/{pid}/match").status_code == 404
    assert client.get(f"/api/projects/{pid}/export").status_code == 404
    assert client.get("/api/projects").json() == []


def test_settings_rates_and_dashboard(client, price):
    assert (
        client.put(
            "/api/settings", json={"exchange_rates": {"EUR:ALL": "100"}}
        ).status_code
        == 200
    )
    assert (
        client.put(
            "/api/settings", json={"exchange_rates": {"EUR:ALL": "0"}}
        ).status_code
        == 422
    )
    data = client.get("/api/dashboard").json()
    assert data["total_prices"] == 1
    assert data["total_values"] == {"ALL": "0", "EUR": "0"}


def test_weak_match_is_persisted_and_exported_as_approved(client, price):
    project = create_project(client)
    pid = project["id"]
    upload_boq(client, project, description="Asfalt rrugor", code="UNKNOWN")
    response = client.post(f"/api/projects/{pid}/match")
    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert row["decision"] == "approved"
    assert row["item_id"] == price["id"]
    assert row["match_result"]["confidence"] == 0
    assert row["match_result"]["classification"] == "low"
    assert row["match_result"]["review_status"] == "approved"
    assert Decimal(row["unit_price"]) == 100
    assert not response.json()["totals"]["incomplete"]
    stored = client.get(f"/api/projects/{pid}").json()["rows"][0]
    for field in ("decision", "item_id", "unit_price", "price_snapshot", "match_result"):
        assert stored[field] == row[field]
    for format in ("standard", "original"):
        export = client.get(f"/api/projects/{pid}/export?format={format}")
        assert export.status_code == 200, export.text
        sheet = load_workbook(BytesIO(export.content)).worksheets[0]
        headers = {cell.value: cell.column for cell in sheet[1]}
        assert sheet.cell(2, headers["Approval status"]).value == "approved"
        confidence = sheet.cell(2, headers["Matching confidence"])
        assert confidence.value == 0
        assert confidence.number_format == "0.00%"
    # Re-running must not teach the matcher that an automatic weak match is strong.
    rerun = client.post(f"/api/projects/{pid}/match").json()["rows"][0]
    assert rerun["match_result"]["confidence"] == 0


def test_no_candidates_remain_unapproved(client, user):
    project = create_project(client)
    upload_boq(client, project)
    pid = project["id"]
    response = client.post(f"/api/projects/{pid}/match")
    assert response.status_code == 200, response.text
    row = response.json()["rows"][0]
    assert row["decision"] == "pending"
    assert row["item_id"] is None
    assert row["match_result"]["explanation"] == "No match found"
    assert row["match_result"]["classification"] == "no_match"
    assert client.get(f"/api/projects/{pid}").json()["rows"][0] == row


def test_automatic_match_can_be_rejected_and_status_stays_consistent(client, price):
    project = create_project(client)
    pid = project["id"]
    upload_boq(client, project)
    row = client.post(f"/api/projects/{pid}/match").json()["rows"][0]
    rejected = client.patch(
        f"/api/projects/{pid}/rows/{row['id']}",
        json={"decision": "rejected", "confirm_override": True},
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["match_result"]["review_status"] == "rejected"
    response = client.post(f"/api/projects/{pid}/match").json()
    assert response["rows"][0]["decision"] == "rejected"
    assert response["rows"][0]["unit_price"] is None
